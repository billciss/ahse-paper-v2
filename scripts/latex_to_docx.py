"""
latex_to_docx.py
================
Converts the manuscript (paper/main.tex and its \\input sections) to a Word document, for
co-authors who work in Word. It handles the LaTeX subset used by this manuscript: front matter
(title, authors, affiliations, abstract, highlights, keywords), sections, paragraphs, lists,
tables (tabular with booktabs, multicolumn), figures (PNG versions), numbered equations, inline
math (rendered as text with sub/superscripts), numbered citations with a reference list built
from references.bib, and cross-references.

Usage (WSL):  python scripts/latex_to_docx.py --paper paper --out paper/manuscript.docx
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# ----------------------------------------------------------------------------- text helpers
SYMBOLS = {
    r"\alpha": "α", r"\beta": "β", r"\gamma": "γ", r"\delta": "δ", r"\eta": "η", r"\tau": "τ",
    r"\sigma": "σ", r"\mu": "μ", r"\Delta": "Δ", r"\le": "≤", r"\leq": "≤", r"\ge": "≥",
    r"\geq": "≥", r"\in": "∈", r"\pm": "±", r"\to": "→", r"\times": "×", r"\approx": "≈",
    r"\lceil": "⌈", r"\rceil": "⌉", r"\dots": "…", r"\ldots": "…", r"\cdot": "·", r"\circ": "°",
    r"\infty": "∞", r"\sim": "∼", r"\neq": "≠", r"\{": "{", r"\}": "}", r"\%": "%", r"\,": "\u2009",
    r"\;": " ", r"\!": "", r"\quad": "  ", r"\qquad": "    ", r"\max": "max", r"\min": "min",
    r"\ln": "ln", r"\exp": "exp", r"\colon": ":", r"\mid": "|",
}
DROP = (r"\left", r"\right", r"\bigl", r"\bigr", r"\Bigl", r"\Bigr", r"\big", r"\Big",
        r"\displaystyle", r"\nolimits", r"\limits")
ACCENTS = {"`": "\u0300", "'": "\u0301", "^": "\u0302", '"': "\u0308", "~": "\u0303"}


def accents(s):
    """{\\`e}, \\'e, {\\L} ... -> unicode."""
    s = re.sub(r"\{\\L\}|\\L\b", "Ł", s)
    s = re.sub(r"\{\\l\}", "ł", s)

    def acc(m):
        return m.group(2) + ACCENTS[m.group(1)]
    s = re.sub(r"\{\\([`'^\"~])\{?(\w)\}?\}", acc, s)
    s = re.sub(r"\\([`'^\"~])\{?(\w)\}?", acc, s)
    import unicodedata
    return unicodedata.normalize("NFC", s)


def read_group(s, i):
    """s[i] == '{' -> (content, index after the closing brace)."""
    assert s[i] == "{", s[i:i + 30]
    depth, j = 0, i
    while j < len(s):
        if s[j] == "\\":
            j += 2
            continue
        if s[j] == "{":
            depth += 1
        elif s[j] == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    raise ValueError("unbalanced braces: " + s[i:i + 60])


def math_runs(m, style=""):
    """Inline math -> list of (text, style) with style in '', 'sub', 'sup'."""
    out, i = [], 0
    m = m.strip()
    while i < len(m):
        c = m[i]
        if c in "_^":
            st = "sub" if c == "_" else "sup"
            i += 1
            while i < len(m) and m[i] == " ":
                i += 1
            if i < len(m) and m[i] == "{":
                g, i = read_group(m, i)
                out += [(t, st) for t, _ in math_runs(g, st)]
            elif m.startswith("\\", i):
                cmd = re.match(r"\\[A-Za-z]+|\\.", m[i:]).group(0)
                i += len(cmd)
                if cmd in (r"\mathrm", r"\text", r"\mathit", r"\mathbf", r"\textrm") and i < len(m) and m[i] == "{":
                    g, i = read_group(m, i)
                    out += [(t, st) for t, _ in math_runs(g, st)]
                else:
                    out.append((SYMBOLS.get(cmd, cmd.lstrip("\\")), st))
            else:
                out.append((m[i], st))
                i += 1
            continue
        if c == "\\":
            mm = re.match(r"\\[A-Za-z]+|\\.", m[i:])
            cmd = mm.group(0)
            i += len(cmd)
            if cmd in DROP:
                if i < len(m) and m[i] in "([|.":
                    out.append((m[i] if m[i] != "." else "", style))
                    i += 1
                continue
            if cmd in (r"\mathrm", r"\text", r"\mathit", r"\mathbf", r"\operatorname", r"\textrm"):
                g, i = read_group(m, i)
                out += math_runs(g.replace(" ", "\u00a0") if cmd == r"\text" else g, style)
                continue
            if cmd == r"\mathcal":
                g, i = read_group(m, i)
                out.append((g, style))
                continue
            if cmd in (r"\hat", r"\bar", r"\tilde"):
                g, i = (read_group(m, i) if m[i] == "{" else (m[i], i + 1))
                mark = {r"\hat": "\u0302", r"\bar": "\u0304", r"\tilde": "\u0303"}[cmd]
                inner = "".join(t for t, _ in math_runs(g, style))
                out.append((inner[:1] + mark + inner[1:], style))
                continue
            if cmd in (r"\frac", r"\tfrac", r"\dfrac"):
                a, i = read_group(m, i)
                b, i = read_group(m, i)
                out += math_runs(a, style) + [("/", style)] + math_runs(b, style)
                continue
            if cmd == r"\sqrt":
                g, i = read_group(m, i)
                out += [("√(", style)] + math_runs(g, style) + [(")", style)]
                continue
            out.append((SYMBOLS.get(cmd, cmd.lstrip("\\")), style))
            continue
        if c == "{":
            g, i = read_group(m, i)
            out += math_runs(g, style)
            continue
        if c == "}":
            i += 1
            continue
        if c == "~":
            out.append(("\u00a0", style))
            i += 1
            continue
        out.append((c, style))
        i += 1
    # merge
    merged = []
    for t, st in out:
        if merged and merged[-1][1] == st:
            merged[-1] = (merged[-1][0] + t, st)
        elif t:
            merged.append((t, st))
    return merged


def plain(s):
    """Text without formatting (titles, captions inside tables, bib fields)."""
    return "".join(t for t, _ in text_runs(s, None))


def text_runs(s, ctx, base=""):
    """LaTeX text -> list of (text, style); style is a '+'-joined set of b/i/sub/sup/url."""
    out, i = [], 0
    s = s.replace("---", "—").replace("--", "–").replace("``", "“").replace("''", "”")
    while i < len(s):
        c = s[i]
        if c == "$":
            j = s.index("$", i + 1)
            out += [(t, "+".join(x for x in (base, st) if x)) for t, st in math_runs(s[i + 1:j])]
            i = j + 1
            continue
        if c == "\\":
            mm = re.match(r"\\[A-Za-z]+\*?|\\.", s[i:])
            cmd = mm.group(0)
            i += len(cmd)
            if cmd in (r"\textbf", r"\emph", r"\textit", r"\textsc", r"\texttt", r"\textrm", r"\mbox"):
                g, i = read_group(s, i)
                st = {"\\textbf": "b", "\\emph": "i", "\\textit": "i"}.get(cmd, "")
                out += text_runs(g, ctx, "+".join(x for x in (base, st) if x))
                continue
            if cmd in (r"\cite", r"\citep", r"\citet"):
                g, i = read_group(s, i)
                nums = [ctx.cite(k.strip()) for k in g.split(",")] if ctx else []
                out.append(("[" + compress(nums) + "]", base))
                continue
            if cmd in (r"\ref", r"\eqref", r"\cref", r"\Cref"):
                g, i = read_group(s, i)
                n = ctx.ref(g) if ctx else "?"
                out.append((f"({n})" if cmd == r"\eqref" else n, base))
                continue
            if cmd == r"\url":
                g, i = read_group(s, i)
                out.append((g, "+".join(x for x in (base, "url") if x)))
                continue
            if cmd == r"\label":
                _, i = read_group(s, i)
                continue
            if cmd in (r"\textasciitilde",):
                out.append(("~", base))
                continue
            if cmd in (r"\%", r"\&", r"\_", r"\#", r"\$"):
                out.append((cmd[1], base))
                continue
            if cmd in (r"\,", r"\ ", r"\@", r"\/"):
                out.append(("\u2009" if cmd == r"\," else (" " if cmd == r"\ " else ""), base))
                continue
            if cmd in (r"\smallskip", r"\medskip", r"\bigskip", r"\centering", r"\small", r"\footnotesize",
                       r"\FloatBarrier", r"\noindent", r"\newline", r"\\"):
                continue
            if cmd in SYMBOLS:
                out.append((SYMBOLS[cmd], base))
                continue
            if i < len(s) and s[i] == "{":      # unknown command with argument: keep the argument
                g, i = read_group(s, i)
                out += text_runs(g, ctx, base)
            continue
        if c == "{":
            g, i = read_group(s, i)
            out += text_runs(g, ctx, base)
            continue
        if c == "}":
            i += 1
            continue
        if c == "~":
            out.append(("\u00a0", base))
            i += 1
            continue
        out.append((c, base))
        i += 1
    merged = []
    for t, st in out:
        t = accents(t) if "\\" in t else t
        if merged and merged[-1][1] == st:
            merged[-1] = (merged[-1][0] + t, st)
        elif t:
            merged.append((t, st))
    return [(re.sub(r"[ \t\r\n]+", " ", t), st) for t, st in merged]


def compress(nums):
    nums = sorted(set(nums))
    parts, k = [], 0
    while k < len(nums):
        j = k
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        parts.append(str(nums[k]) if j == k else (f"{nums[k]},{nums[j]}" if j == k + 1 else f"{nums[k]}–{nums[j]}"))
        k = j + 1
    return ",".join(parts)


# ----------------------------------------------------------------------------- bibliography
def parse_bib(path):
    text = Path(path).read_text(encoding="utf-8")
    entries = {}
    for m in re.finditer(r"@(\w+)\s*\{\s*([^,\s]+)\s*,", text):
        start = m.end()
        depth, j = 1, m.end() - 1
        j = text.index("{", m.start())
        _, end = read_group(text, j)
        body = text[start:end - 1]
        fields = {}
        for f in re.finditer(r"(\w+)\s*=\s*", body):
            k = f.group(1).lower()
            p = f.end()
            if p < len(body) and body[p] == "{":
                v, _ = read_group(body, p)
            elif p < len(body) and body[p] == '"':
                v = body[p + 1:body.index('"', p + 1)]
            else:
                v = re.match(r"[^,\n]+", body[p:]).group(0)
            fields.setdefault(k, v.strip())
        entries[m.group(2)] = (m.group(1).lower(), fields)
    return entries


def fmt_authors(a):
    names = [n.strip() for n in re.split(r"\s+and\s+", a) if n.strip()]
    out = []
    for n in names:
        if n.startswith("{") and n.endswith("}"):       # corporate author, e.g. {IRENA}
            out.append(plain(n))
            continue
        n = plain(n)
        if "," not in n:
            if " " in n and not n.startswith("National"):
                parts = n.split()
                out.append(" ".join(p[0] + "." for p in parts[:-1]) + " " + parts[-1])
            else:
                out.append(n.strip("{}"))
            continue
        last, first = [x.strip() for x in n.split(",", 1)]
        ini = " ".join(w[0] + "." for w in re.split(r"[\s]+", first) if w)
        ini = re.sub(r"(\w)\.-(\w)", r"\1.-\2", ini)
        out.append(f"{ini} {last}")
    if len(out) > 6:
        out = out[:6] + ["et al."]
        return ", ".join(out[:-1]) + " " + out[-1]
    return ", ".join(out)


def fmt_entry(kind, f):
    g = lambda k: plain(f[k]) if k in f else ""  # noqa: E731
    parts = []
    if "author" in f:
        parts.append(fmt_authors(f["author"]) + ",")
    parts.append(g("title").rstrip(".") + ",")
    if kind == "article":
        src = g("journal") + (f" {g('volume')}" if "volume" in f else "")
        src += f" ({g('year')})" if "year" in f else ""
        src += f" {g('pages')}" if "pages" in f else ""
        parts.append(src + ".")
    elif kind in ("inproceedings", "incollection"):
        parts.append("in: " + g("booktitle") + (f", vol. {g('volume')}" if "volume" in f else "")
                     + (f", pp. {g('pages')}" if "pages" in f else "") + f", {g('year')}.")
    elif kind == "book":
        parts.append(f"{g('publisher')}, {g('address') + ', ' if 'address' in f else ''}{g('year')}.")
    else:
        extra = ", ".join(x for x in (g("howpublished"), g("institution"), g("publisher"), g("note")) if x)
        parts.append((extra + ", " if extra else "") + g("year") + ".")
    if "doi" in f:
        parts.append("https://doi.org/" + f["doi"].strip())
    elif "url" in f:
        parts.append(f["url"].strip())
    return re.sub(r"\s+", " ", " ".join(parts)).replace(",,", ",").replace(" ,", ",")


# ----------------------------------------------------------------------------- document
class Ctx:
    def __init__(self, bib):
        self.bib, self.order, self.labels = bib, [], {}

    def cite(self, k):
        if k not in self.order:
            self.order.append(k)
        return self.order.index(k) + 1

    def ref(self, label):
        return str(self.labels.get(label, "??"))


def strip_comments(t):
    return "\n".join(re.sub(r"(?<!\\)%.*", "", line) for line in t.splitlines())


def add_runs(p, runs, size=None):
    for t, st in runs:
        r = p.add_run(t)
        if "b" in st.split("+"):
            r.bold = True
        if "i" in st.split("+"):
            r.italic = True
        if "sub" in st.split("+"):
            r.font.subscript = True
        if "sup" in st.split("+"):
            r.font.superscript = True
        if size:
            r.font.size = Pt(size)
    return p


def collect_labels(body, ctx):
    """Number sections, figures, tables and equations in document order."""
    sec = sub = fig = tab = eq = 0
    for m in re.finditer(r"\\(section|subsection)\*?\{|\\begin\{(figure|table|equation)\}|\\label\{([^}]+)\}", body):
        if m.group(1) == "section" and not body[m.start():m.end()].startswith(r"\section*"):
            sec += 1
            sub = 0
            cur = str(sec)
        elif m.group(1) == "subsection" and not body[m.start():m.end()].startswith(r"\subsection*"):
            sub += 1
            cur = f"{sec}.{sub}"
        elif m.group(2) == "figure":
            fig += 1
            cur = str(fig)
        elif m.group(2) == "table":
            tab += 1
            cur = str(tab)
        elif m.group(2) == "equation":
            eq += 1
            cur = str(eq)
        elif m.group(3):
            ctx.labels[m.group(3)] = cur


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_table(doc, block, ctx, number):
    cap = re.search(r"\\caption\{", block)
    caption, _ = read_group(block, cap.end() - 1) if cap else ("", 0)
    p = doc.add_paragraph()
    add_runs(p, [(f"Table {number}. ", "b")] + text_runs(caption, ctx), 9)
    tab_start = block.index(r"\begin{tabular}")
    spec, k = read_group(block, block.index("{", tab_start + len(r"\begin{tabular}")))
    body = block[k:block.index(r"\end{tabular}")]
    rows = []
    for line in re.split(r"\\\\", body):
        line = re.sub(r"\\(toprule|midrule|bottomrule|hline)", "", line)
        line = re.sub(r"\\cmidrule(\([^)]*\))?\{[^}]*\}", "", line).strip()
        if not line:
            continue
        cells = []
        for c in re.split(r"(?<!\\)&", line):
            c = c.strip()
            mc = re.match(r"\\multicolumn\{(\d+)\}\{[^}]*\}\{", c)
            if mc:
                inner, _ = read_group(c, mc.end() - 1)
                cells.append((inner, int(mc.group(1))))
            else:
                cells.append((c, 1))
        rows.append(cells)
    ncol = max(sum(w for _, w in r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, r in enumerate(rows):
        ci = 0
        for text, w in r:
            cell = t.cell(ri, ci)
            if w > 1:
                cell = cell.merge(t.cell(ri, ci + w - 1))
            cell.text = ""
            add_runs(cell.paragraphs[0], text_runs(text, ctx), 8)
            if ri == 0 or (ri == 1 and any(w > 1 for _, w in rows[0])):
                set_cell_shading(cell, "E7EEF7")
            ci += w
    note = block[block.index(r"\end{tabular}") + len(r"\end{tabular}"):]
    note = re.sub(r"\\caption\{.*", "", note, flags=re.S)
    note = re.sub(r"\\label\{[^}]*\}|\\(smallskip|footnotesize|centering|small)", "", note).strip()
    if note:
        add_runs(doc.add_paragraph(), text_runs(note, ctx), 8)
    doc.add_paragraph()


def add_figure(doc, block, ctx, number, figdir):
    g = re.search(r"\\includegraphics(\[[^]]*\])?\{([^}]+)\}", block).group(2)
    png = figdir / (Path(g).stem + ".png")
    if png.exists():
        doc.add_picture(str(png), width=Cm(16))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = re.search(r"\\caption\{", block)
    caption, _ = read_group(block, cap.end() - 1)
    p = doc.add_paragraph()
    add_runs(p, [(f"Fig. {number}. ", "b")] + text_runs(caption, ctx), 9)


def add_equation(doc, math, number):
    math = re.sub(r"\\label\{[^}]*\}", "", math)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, [(t, ("i+" + st).strip("+")) for t, st in math_runs(math)])
    p.add_run(f"\t\t({number})")


def add_list(doc, block, ctx, numbered):
    items = re.split(r"\\item\b", block)[1:]
    for n, it in enumerate(items, 1):
        p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
        add_runs(p, text_runs(it.strip(), ctx))


def render_body(doc, body, ctx, figdir):
    counters = {"figure": 0, "table": 0, "equation": 0}
    sec = sub = 0
    pos = 0
    pattern = re.compile(r"\\(section|subsection|paragraph)(\*?)\{|\\begin\{(figure|table|equation|enumerate|itemize)\}(\[[^]]*\])?")
    buf = ""

    def flush(text):
        for para in re.split(r"\n\s*\n", text):
            para = para.strip()
            if para and re.sub(r"\\(label|FloatBarrier)\{?[^}]*\}?", "", para).strip():
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                add_runs(p, text_runs(para, ctx))

    for m in pattern.finditer(body):
        if m.start() < pos:
            continue
        flush(body[pos:m.start()])
        if m.group(1):
            title, end = read_group(body, m.end() - 1)
            star = m.group(2) == "*"
            if m.group(1) == "section":
                sec, sub = (sec if star else sec + 1), 0
                doc.add_heading(("" if star else f"{sec}. ") + plain(title), level=1)
            elif m.group(1) == "subsection":
                sub = sub if star else sub + 1
                doc.add_heading(("" if star else f"{sec}.{sub}. ") + plain(title), level=2)
            else:
                rest_end = body.find("\n\n", end)
                rest_end = len(body) if rest_end < 0 else rest_end
                p = doc.add_paragraph()
                add_runs(p, [(plain(title) + " ", "b+i")] + text_runs(body[end:rest_end], ctx))
                end = rest_end
            pos = end
            continue
        env = m.group(3)
        close = r"\end{%s}" % env
        end = body.index(close, m.end()) + len(close)
        block = body[m.end():end - len(close)]
        if env in ("figure", "table", "equation"):
            counters[env] += 1
        if env == "figure":
            add_figure(doc, block, ctx, counters[env], figdir)
        elif env == "table":
            add_table(doc, block, ctx, counters[env])
        elif env == "equation":
            add_equation(doc, block, counters[env])
        else:
            add_list(doc, block, ctx, env == "enumerate")
        pos = end
    flush(body[pos:])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--paper", type=Path, default=Path("paper"))
    ap.add_argument("--out", type=Path, default=Path("paper/manuscript.docx"))
    args = ap.parse_args()
    root = args.paper
    main_tex = strip_comments((root / "main.tex").read_text(encoding="utf-8"))
    body_start = main_tex.index(r"\end{frontmatter}")
    front, rest = main_tex[:body_start], main_tex[body_start:]
    rest = re.sub(r"\\input\{([^}]+)\}",
                  lambda m: strip_comments((root / (m.group(1) + ".tex")).read_text(encoding="utf-8")), rest)
    rest = rest[:rest.index(r"\bibliographystyle")]
    rest = rest.replace(r"\end{frontmatter}", "")

    bib = parse_bib(root / "references.bib")
    ctx = Ctx(bib)
    collect_labels(rest, ctx)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.5)

    # ---- front matter
    title, _ = read_group(front, front.index(r"\title{") + len(r"\title"))
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(h, [(plain(title), "b")], 16)
    affs = {m.group(1): m for m in re.finditer(r"\\affiliation\[(\d+)\]\{", front)}
    authors = []
    for m in re.finditer(r"\\author\[([^]]*)\]\{", front):
        name, _ = read_group(front, m.end() - 1)
        corr = r"\corref" in name
        name = re.sub(r"\\corref\{[^}]*\}", "", name).strip()
        authors.append((name, m.group(1), corr))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for k, (name, marks, corr) in enumerate(authors):
        add_runs(p, [(accents(name), ""), (marks.replace(",", ",") + ("*" if corr else ""), "sup")])
        if k < len(authors) - 1:
            p.add_run(", ")
    for num, m in affs.items():
        g, _ = read_group(front, m.end() - 1)
        fields = dict(re.findall(r"(\w+)=\{([^}]*(?:\{[^}]*\}[^}]*)*)\}", g))
        txt = ", ".join(re.sub(r"\s+", " ", fields[k]) for k in ("organization", "city", "postcode", "country") if k in fields)
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(q, [(num, "sup"), (" " + txt, "i")], 9)
    emails = re.findall(r"\\ead\{([^}]*)\}", front)
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(q, [("* Corresponding author. E-mail addresses: " + ", ".join(emails), "")], 9)

    abstract = front[front.index(r"\begin{abstract}") + len(r"\begin{abstract}"):front.index(r"\end{abstract}")]
    doc.add_heading("Abstract", level=1)
    add_runs(doc.add_paragraph(), text_runs(abstract.strip(), ctx))
    if r"\begin{highlights}" in front:
        doc.add_heading("Highlights", level=1)
        hl = front[front.index(r"\begin{highlights}") + 18:front.index(r"\end{highlights}")]
        add_list(doc, hl, ctx, False)
    if r"\begin{keyword}" in front:
        kw = front[front.index(r"\begin{keyword}") + 15:front.index(r"\end{keyword}")]
        p = doc.add_paragraph()
        add_runs(p, [("Keywords: ", "b")] + text_runs("; ".join(x.strip() for x in kw.split(r"\sep")), ctx))
    if r"\begin{graphicalabstract}" in front:
        ga = front[front.index(r"\begin{graphicalabstract}"):front.index(r"\end{graphicalabstract}")]
        g = re.search(r"\\includegraphics(\[[^]]*\])?\{([^}]+)\}", ga).group(2)
        png = root / "figures" / (Path(g).stem + ".png")
        if png.exists():
            doc.add_heading("Graphical abstract", level=1)
            doc.add_picture(str(png), width=Cm(16))

    # ---- body
    render_body(doc, rest, ctx, root / "figures")

    # ---- references
    doc.add_heading("References", level=1)
    for n, k in enumerate(ctx.order, 1):
        kind, f = bib[k]
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        add_runs(p, [(f"[{n}] " + fmt_entry(kind, f), "")], 9)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    missing = [k for k in ctx.labels.values() if k == "??"]
    print(f"written {args.out}: {len(ctx.order)} references, {len(ctx.labels)} labels")


if __name__ == "__main__":
    main()
